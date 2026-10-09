"""Esporta la tabella dei settori in un file Excel formattato."""

from __future__ import annotations

import io
from collections.abc import Callable
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PCT, PCT2, NUM, INT = "0.0%", "0.00%", "0.00", "#,##0"

# (gruppo, intestazione, estrattore, formato)
Col = tuple[str, str, Callable[[dict], Any], str | None]
COLUMNS: list[Col] = [
    ("", "Settore", lambda r: r["name"], None),
    ("", "ETF", lambda r: r["symbol"], None),
    ("", "Prezzo", lambda r: r["price"], NUM),
    ("", "Ultimo dato", lambda r: r["last_date"], None),
    ("Trend", "Stato", lambda r: r["trend_label"], None),
    ("Trend", "vs SMA50", lambda r: r["px_vs_sma50"], PCT),
    ("Trend", "vs SMA200", lambda r: r["px_vs_sma200"], PCT),
    ("Trend", "SMA50 vs SMA200", lambda r: r["sma50_vs_sma200"], PCT),
    ("Trend", "Score", lambda r: r["scores"]["trend"], "score"),
    ("Momentum", "1m", lambda r: r["ret_1m"], PCT),
    ("Momentum", "3m", lambda r: r["ret_3m"], PCT),
    ("Momentum", "6m", lambda r: r["ret_6m"], PCT),
    ("Momentum", "12m", lambda r: r["ret_12m"], PCT),
    ("Momentum", "RSI 14", lambda r: r["rsi14"], "0"),
    ("Momentum", "RS 3m vs SPY", lambda r: r["rs_3m"], PCT),
    ("Momentum", "RS 6m vs SPY", lambda r: r["rs_6m"], PCT),
    ("Momentum", "Score", lambda r: r["scores"]["momentum"], "score"),
    ("Volume", "Vol 20g/90g", lambda r: r["vol_ratio_20_90"], NUM),
    ("Volume", "Su/Giù", lambda r: r["updown_volume"], NUM),
    ("Volume", "Controvalore medio 20g", lambda r: r["dollar_volume_20d"], INT),
    ("Volume", "Score", lambda r: r["scores"]["volume"], "score"),
    ("Volatilità", "Vol 20g", lambda r: r["vol_20d"], PCT),
    ("Volatilità", "Vol 60g", lambda r: r["vol_60d"], PCT),
    ("Volatilità", "ATR%", lambda r: r["atr_pct"], PCT2),
    ("Volatilità", "Da max 52w", lambda r: r["drawdown_52w"], PCT),
    ("Volatilità", "Score", lambda r: r["scores"]["volatility"], "score"),
    ("", "Score totale", lambda r: r["scores"]["total"], "score"),
]

NOTES = [
    "Punteggi 0-100: posizione del settore rispetto agli altri 10 (percentile). Non sono segnali operativi.",
    "Trend: prezzo vs medie a 50 e 200 giorni. Rialzista = prezzo > SMA50 > SMA200; Ribassista = il contrario.",
    "Momentum: rendimenti a 1/3/6/12 mesi, RSI(14) e forza relativa (RS) rispetto a SPY.",
    "Volume: volume medio 20g / 90g e rapporto tra volume nei giorni in rialzo e in ribasso (ultimi 20g).",
    "Volatilità: più bassa = punteggio più alto (vol 60g, ATR, distanza dal massimo a 52 settimane).",
    "Feed IEX (gratuito): i volumi sono solo quelli della borsa IEX; i rapporti sono confrontabili, i valori assoluti no.",
]


def _fill(score: float) -> PatternFill:
    """Rosso (0) -> giallo (50) -> verde (100), toni chiari."""
    t = max(0.0, min(100.0, score)) / 100
    r, g = (255, int(120 + 270 * t)) if t < 0.5 else (int(255 - 280 * (t - 0.5)), 205)
    return PatternFill("solid", fgColor=f"{r:02X}{min(g, 215):02X}7A")


def _sheet(
    ws, columns: list[Col], rows: list[dict], first_data_row: int = 3, freeze: str = "C3"
) -> None:
    """Scrive intestazioni di gruppo (riga 1), intestazioni (riga 2) e righe dati."""
    thin = Side(style="thin", color="BBBBBB")
    n = len(columns)
    start = 1
    for i, (grp, head, _, _) in enumerate(columns, 1):
        ws.cell(2, i, head)
        if grp and (i == n or columns[i][0] != grp):
            ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=i)
            ws.cell(1, start, grp)
        if i < n and columns[i][0] != grp:
            start = i + 1
    for c in range(1, n + 1):
        for r_ in (1, 2):
            cell = ws.cell(r_, c)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if r_ == 2:
                cell.border = Border(bottom=thin)
                cell.fill = PatternFill("solid", fgColor="E8ECF4")

    for ri, row in enumerate(rows, first_data_row):
        for ci, (_, head, get, fmt) in enumerate(columns, 1):
            v = get(row)
            cell = ws.cell(ri, ci, v)
            if fmt == "score":
                cell.number_format = "0"
                if v is not None:
                    cell.fill = _fill(v)
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")
            elif fmt:
                cell.number_format = fmt
            if head == "Stato":
                color = {"Rialzista": "12805C", "Ribassista": "C0392B"}.get(v, "444444")
                cell.font = Font(bold=True, color=color)
                cell.alignment = Alignment(horizontal="center")
            if fmt in (PCT, PCT2) and isinstance(v, (int, float)) and v < 0:
                cell.font = Font(color="C0392B")

    for i, (_, head, _, _) in enumerate(columns, 1):
        wide = head in ("Settore", "Azienda")
        ws.column_dimensions[get_column_letter(i)].width = (
            28 if wide else max(10, min(22, len(head) + 3))
        )
    ws.row_dimensions[2].height = 32
    ws.freeze_panes = freeze


def _notes_sheet(wb: Workbook, title: str, lines: list[str]) -> None:
    info = wb.create_sheet("Note")
    info["A1"] = title
    info["A1"].font = Font(bold=True, size=13)
    for i, line in enumerate(lines, 3):
        info.cell(i, 1, line)
    info.column_dimensions["A"].width = 130


def _save(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build(payload: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Settori"
    rows = sorted(payload["sectors"], key=lambda r: r["scores"]["total"] or -1, reverse=True)
    _sheet(ws, COLUMNS, rows)
    lines = [f"Aggiornato: {payload['updated']}"]
    lines.append(f"Feed dati: {payload['feed']}" + (" (DATI SINTETICI)" if payload["demo"] else ""))
    _notes_sheet(wb, "Sector Lab - ETF settoriali SPDR", [*lines, "", *NOTES])
    return _save(wb)


STOCK_COLUMNS: list[Col] = [
    ("", "Pos.", lambda r: r["rank"], "0"),
    ("", "Ticker", lambda r: r["symbol"], None),
    ("", "Azienda", lambda r: r["name"], None),
    ("", "Prezzo", lambda r: r["price"], NUM),
    ("", "Peso in ACWI %", lambda r: r["weight_pct"], "0.000"),
    ("Trend (rel. ACWI)", "Stato", lambda r: r["trend_label"], None),
    ("Trend (rel. ACWI)", "vs SMA50", lambda r: r["rp_vs_sma50"], PCT),
    ("Trend (rel. ACWI)", "vs SMA200", lambda r: r["rp_vs_sma200"], PCT),
    ("Trend (rel. ACWI)", "SMA50 vs SMA200", lambda r: r["rp_sma50_vs_sma200"], PCT),
    ("Trend (rel. ACWI)", "Score", lambda r: r["scores"]["trend"], "score"),
    ("Momentum (TWRR vs ACWI)", "1m", lambda r: r["twrr_1m"], PCT),
    ("Momentum (TWRR vs ACWI)", "3m", lambda r: r["twrr_3m"], PCT),
    ("Momentum (TWRR vs ACWI)", "6m", lambda r: r["twrr_6m"], PCT),
    ("Momentum (TWRR vs ACWI)", "12m", lambda r: r["twrr_12m"], PCT),
    ("Momentum (TWRR vs ACWI)", "TWRR ponderato", lambda r: r["twrr_w"], PCT),
    ("Momentum (TWRR vs ACWI)", "Score", lambda r: r["scores"]["momentum"], "score"),
    ("Volume (rel. ACWI)", "Vol 20g/90g relativo", lambda r: r["vol_ratio_rel"], NUM),
    ("Volume (rel. ACWI)", "Su/Giù relativo", lambda r: r["updown_rel"], NUM),
    ("Volume (rel. ACWI)", "Score", lambda r: r["scores"]["volume"], "score"),
    ("Volatilità (rel. ACWI)", "Tracking error 60g", lambda r: r["te_60d"], PCT),
    ("Volatilità (rel. ACWI)", "Da max relativo 52w", lambda r: r["rel_drawdown"], PCT),
    ("Volatilità (rel. ACWI)", "Vol titolo 60g", lambda r: r["vol_60d"], PCT),
    ("Volatilità (rel. ACWI)", "Score", lambda r: r["scores"]["volatility"], "score"),
    ("Beta (vs ACWI)", "Beta 1 anno", lambda r: r["beta_1y"], NUM),
    ("Beta (vs ACWI)", "Beta 6 mesi", lambda r: r["beta_6m"], NUM),
    ("Beta (vs ACWI)", "Correlazione 1a", lambda r: r["corr_1y"], NUM),
    ("Pesi", "Score indice settore", lambda r: r.get("sector_score"), "score"),
    ("Pesi", "Peso nel settore", lambda r: r.get("weight_in_sector"), PCT),
    ("Pesi", "Peso nel portafoglio", lambda r: r.get("weight_portfolio"), PCT),
    ("", "Score totale", lambda r: r["scores"]["total"], "score"),
]
STOCK_NOTES = [
    "Universo: titoli USA del fondo iShares MSCI ACWI, classificati per settore (11 settori GICS).",
    "Benchmark: ACWI. Tutte le caratteristiche sono calcolate sul prezzo relativo = prezzo titolo / prezzo ACWI.",
    "TWRR (time-weighted relative return) = (1 + rendimento titolo) / (1 + rendimento ACWI) - 1 sulla stessa finestra.",
    "TWRR ponderato = media delle finestre 1m/3m/6m/12m con pesi 4/3/2/1 (le finestre recenti contano di più).",
    "Trend: prezzo relativo vs sue medie a 50 e 200 giorni. Rialzista = relativo > SMA50 > SMA200.",
    "Volume: volume 20g/90g del titolo diviso lo stesso rapporto di ACWI; Su/Giù = volume nei giorni in cui il relativo sale / scende.",
    "Volatilità: tracking error a 60g (dev. standard annualizzata dei rendimenti relativi), distanza dal massimo relativo a 52 settimane, volatilità del titolo. Più bassa = punteggio più alto.",
    "Beta = cov(rend. titolo, rend. ACWI) / var(rend. ACWI) sui rendimenti giornalieri (1 anno e 6 mesi).",
    "Score indice settore = punteggio totale dell'ETF del settore (tabella Settori), confrontato con quello degli altri 10 settori.",
    "Peso nel settore = quota del titolo in proporzione al suo punteggio (solo titoli idonei), senza tetti.",
    (
        "Peso nel portafoglio = peso reale dell'algoritmo di ribilanciamento: il punteggio del titolo è "
        "valutato insieme ai punteggi degli indici degli altri settori (peso del settore x quota nel settore), con tetto "
        "10% per titolo e 97% investito."
    ),
    "Punteggi 0-100 = percentile dentro il settore. Classifica relativa, non un segnale operativo.",
    "Feed IEX (gratuito): i volumi sono solo quelli della borsa IEX; i rapporti sono confrontabili, i valori assoluti no.",
]

BETA_NOTES = [
    "STRATEGIA ALTO BETA (sistematica, solo acquisto): per ogni settore",
    "  1. Filtro: trend relativo non Ribassista e TWRR ponderato > 0 (il titolo sta battendo ACWI).",
    "  2. Priorità: tra gli idonei, i 10 con beta a 1 anno più alto.",
    "  3. Il beta sceglie i titoli; il peso nel settore segue il punteggio, senza tetto per titolo.",
    "  Se gli idonei sono meno di 10 la lista è più corta: nessun titolo non idoneo viene aggiunto.",
    "  Un beta alto amplifica sia i guadagni sia le perdite rispetto al mercato: è una scelta di rischio, non di qualità.",
]


def build_stocks(
    results: dict[str, dict], sector_names: dict[str, str], meta: dict, mode: str = "quality"
) -> bytes:
    """results: ETF -> {"top": [...], ...}. Un foglio per settore + foglio riepilogo se più settori."""
    wb = Workbook()
    wb.remove(wb.active)
    cols = STOCK_COLUMNS
    if len(results) > 1:
        ws = wb.create_sheet("Tutti")
        all_rows = []
        for etf, res in results.items():
            all_rows += [{**r, "sector": sector_names[etf]} for r in res["top"]]
        _sheet(ws, [("", "Settore", lambda r: r["sector"], None), *cols], all_rows, freeze="D3")
    for etf, res in results.items():
        ws = wb.create_sheet(f"{etf} {sector_names[etf]}"[:31])
        _sheet(ws, cols, res["top"], freeze="D3")
    lines = [f"Aggiornato: {meta['updated']}"]
    lines.append(f"Feed dati: {meta['feed']}" + (" (DATI SINTETICI)" if meta["demo"] else ""))
    for etf, res in results.items():
        lines.append(
            f"{etf} {sector_names[etf]}: analizzati {res['n_analyzed']}, esclusi {res['n_skipped']} (storico < 1 anno o dati mancanti)"
        )
    title = (
        "Sector Lab - strategia alto beta per settore"
        if mode == "beta"
        else "Sector Lab - migliori aziende per settore"
    )
    _notes_sheet(wb, title, [*lines, "", *(BETA_NOTES if mode == "beta" else []), *STOCK_NOTES])
    return _save(wb)


def build_journal(events: list[dict]) -> bytes:
    """Journal completo (dal più vecchio al più recente) con il dettaglio JSON di ogni evento."""
    import json

    from journal import STRATEGY_LABEL, summarize

    wb = Workbook()
    ws = wb.active
    ws.title = "Journal"
    cols: list[Col] = [
        ("", "N.", lambda e: e["id"], "0"),
        ("", "Data/ora (UTC)", lambda e: e["ts"], None),
        ("", "Portafoglio", lambda e: e.get("portfolio"), None),
        ("", "Strategia", lambda e: STRATEGY_LABEL.get(e.get("strategy"), e.get("strategy")), None),
        ("", "Evento", lambda e: e["kind"], None),
        ("", "Esecuzione", lambda e: e["run_id"], None),
        ("", "Proposta", lambda e: e["proposal_id"], None),
        ("", "Riepilogo", lambda e: summarize(e["kind"], e["payload"]), None),
        (
            "",
            "Dettaglio (JSON)",
            lambda e: json.dumps(e["payload"], ensure_ascii=False, default=str)[:32000],
            None,
        ),
    ]
    _sheet(ws, cols, sorted(events, key=lambda e: e["id"]), freeze="A3")
    for col, width in {"B": 22, "C": 18, "D": 12, "F": 24, "G": 14, "H": 90, "I": 60}.items():
        ws.column_dimensions[col].width = width
    return _save(wb)
