import io
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))
import compare_portfolios as cp
import journal as journal_mod


def _res():
    ba, bb = cp._demo_brokers()
    a, b = cp.snapshot(ba, "Portafoglio 17", "1M"), cp.snapshot(bb, "Portafoglio 18", "1M")
    return cp.compare(a, b), a, b


def test_read_env_ignores_comments_quotes_and_bom(tmp_path):
    f = tmp_path / ".env.x"
    f.write_text(
        '﻿# commento\nALPACA_API_KEY="abc"\nALPACA_SECRET_KEY = s3cr3t \n\nALTRO=1\n',
        encoding="utf-8",
    )
    assert cp.read_env(f)["ALPACA_API_KEY"] == "abc" and cp.credentials(f) == ("abc", "s3cr3t")
    g = tmp_path / ".env.vuoto"
    g.write_text("ALPACA_API_KEY=solo\n")
    with pytest.raises(SystemExit):
        cp.credentials(g)


def test_overlap_sectors_and_weights_are_consistent():
    res, a, b = _res()
    o = res["overlap"]
    wa = {p["symbol"]: p["weight"] for p in a["positions"]}
    wb = {p["symbol"]: p["weight"] for p in b["positions"]}
    assert (
        o["n_a"] == len(wa) and o["n_b"] == len(wb) and o["n_common"] == len(set(wa) & set(wb)) > 0
    )
    assert o["weight_overlap"] == pytest.approx(
        sum(min(wa.get(s, 0), wb.get(s, 0)) for s in set(wa) | set(wb))
    )
    assert o["jaccard"] == pytest.approx(o["n_common"] / len(set(wa) | set(wb)))
    assert sum(d["w_a"] for d in res["sectors"]) == pytest.approx(sum(wa.values()))
    assert sum(d["w_b"] for d in res["sectors"]) == pytest.approx(sum(wb.values()))
    assert all(r["diff"] == pytest.approx(r["w_a"] - r["w_b"]) for r in res["positions"])
    assert {r["in"] for r in res["positions"]} == {"entrambi", "Portafoglio 17", "Portafoglio 18"}


def test_window_metrics_use_only_common_days_and_curve_is_base_100():
    res, a, _ = _res()
    w = res["window"]
    s = a["history"]
    assert w["a"]["total_return"] == pytest.approx(s.iloc[-1] / s.iloc[0] - 1)
    assert w["a"]["max_drawdown"] <= 0 and w["a"]["worst_day"] <= w["a"]["best_day"]
    cur = w["curve"]
    assert list(cur.columns)[:2] == ["Portafoglio 17", "Portafoglio 18"] and cur.iloc[0].tolist()[
        :2
    ] == [100.0, 100.0]


def test_metrics_edge_cases():
    import pandas as pd

    m = cp.series_metrics(pd.Series([100.0]))
    assert m["total_return"] is None and m["n_days"] == 1
    m = cp.series_metrics(pd.Series([100.0, 110.0, 99.0, 120.0]))
    assert m["total_return"] == pytest.approx(0.2) and m["max_drawdown"] == pytest.approx(
        99 / 110 - 1
    )


def test_report_and_excel_contain_both_portfolios(tmp_path):
    res, _, _ = _res()
    ja = {"strategy": "beta", "n_runs": 3, "trade_failed": 1}
    jb = {"strategy": "quality", "n_runs": 2, "trade_failed": 0}
    txt = cp.report(res, ja, jb)
    assert (
        "Portafoglio 17" in txt
        and "Portafoglio 18" in txt
        and "Alto beta" in txt
        and "Qualità" in txt
    )
    assert "Titoli in comune" in txt and "Finestra comune" in txt
    out = tmp_path / "c.xlsx"
    cp.to_excel(res, ja, jb, out)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Riepilogo", "Posizioni", "Settori", "Andamento (base 100)", "Note"]
    rows = {r[0]: r for r in wb["Riepilogo"].iter_rows(min_row=3, values_only=True)}
    assert rows["Strategia (journal)"][1:3] == ("Alto beta", "Qualità") and rows[
        "Ribilanciamenti (journal)"
    ][1:3] == (3, 2)
    assert wb["Posizioni"].max_row > 20


def test_journal_summary_reads_strategy_read_only(tmp_path):
    path = tmp_path / "j.db"
    j = journal_mod.Journal(path, "Portafoglio 18", "quality")
    j.log("RUN", {"params": {"mode": "quality"}, "account": {"equity": 1_000_000.0}}, "r1")
    j.log("TRADE_FAILED", {"symbol": "A", "reason": "x"}, "r1")
    before = path.read_bytes()
    s = cp.journal_summary(path)
    assert s["n_runs"] == 1 and s["strategy"] == "quality" and s["portfolio"] == "Portafoglio 18"
    assert s["trade_failed"] == 1 and s["first_equity"] == 1_000_000.0
    assert path.read_bytes() == before  # sola lettura
    assert cp.journal_summary(tmp_path / "manca.db") == {} and cp.journal_summary(None) == {}


def test_cli_demo_writes_the_excel_and_requires_two_envs(tmp_path, capsys):
    out = tmp_path / "out.xlsx"
    assert cp.main(["--demo", "--out", str(out), "--name-a", "A", "--name-b", "B"]) == 0
    assert out.exists() and "Excel scritto" in capsys.readouterr().out
    assert load_workbook(io.BytesIO(out.read_bytes())).sheetnames[0] == "Riepilogo"
    with pytest.raises(SystemExit):
        cp.main([])  # senza --demo servono entrambi i file di chiavi


def test_window_can_start_from_the_most_recent_portfolio_and_curve_restarts_at_100():
    import pandas as pd

    ba, bb = cp._demo_brokers()
    a, b = cp.snapshot(ba, "A", "1M"), cp.snapshot(bb, "B", "1M")
    start = a["history"].index[8]
    res = cp.compare(a, b, start=start, own_starts=(start, None))
    w = res["window"]
    assert w["from"] == cp._fmt_ts(start) and w["a"]["n_points"] == len(a["history"]) - 8
    sa = a["history"].loc[start:]
    assert w["a"]["total_return"] == pytest.approx(sa.iloc[-1] / sa.iloc[0] - 1)
    assert w["curve"].iloc[0].tolist()[:2] == [100.0, 100.0]
    assert res["own"]["a"]["n_points"] == len(sa) and res["own"]["b"]["n_points"] == len(
        b["history"]
    )
    assert cp.compare(a, b, start=a["history"].index[-1] + pd.Timedelta(days=5))["window"] == {}


def test_volatility_is_annualized_with_the_bars_per_year_of_the_timeframe():
    import pandas as pd

    s = pd.Series([100.0, 100.5, 100.1, 100.8, 100.4, 101.0])
    daily = cp.series_metrics(s, 252)["vol"]
    m15 = cp.series_metrics(s, 252 * cp.BARS_PER_DAY["15Min"])["vol"]
    assert m15 == pytest.approx(daily * cp.BARS_PER_DAY["15Min"] ** 0.5)
    assert cp.BARS_PER_DAY["1D"] == 1 and cp.MINUTES["1H"] == 60


def test_timestamp_formatting_for_daily_and_intraday_points():
    import pandas as pd

    assert cp._fmt_ts(pd.Timestamp("2026-10-09")) == "2026-10-09"
    assert cp._fmt_ts(pd.Timestamp("2026-10-09 14:30")) == "2026-10-09 14:30"


def test_journal_summary_reports_the_first_filled_order(tmp_path):
    j = journal_mod.Journal(tmp_path / "f.db", "P", "short")
    j.log("RUN", {"params": {"mode": "short"}, "account": {"equity": 1.0}}, "r1")
    for pid, status in (("p1", "failed"), ("p2", "filled")):
        j.add_proposal("r1", {"id": pid, "kind": "BUY", "side": "buy", "symbol": "A", "name": "a", "sector": "XLK",
                              "qty": 1, "price": 1.0, "value": 1.0, "reason": "r", "priority": 2})  # fmt: skip
        j.set_status(pid, status, "DECISION_DONE")
    s = cp.journal_summary(tmp_path / "f.db")
    assert s["first_fill"] and s["first_fill"].endswith("Z")
    assert cp.journal_summary(None) == {}


def test_alpaca_history_is_dated_for_1D_and_timestamped_for_intraday():
    from types import SimpleNamespace

    from broker import AlpacaBroker

    ts_day = 1_791_504_000  # 2026-10-09 00:00 UTC
    fake = SimpleNamespace(
        get_portfolio_history=lambda req: SimpleNamespace(
            timestamp=[ts_day, ts_day + 52_200, ts_day + 53_100],
            equity=[1_000_000.0, None, 1_001_000.0],
        )
    )
    br = AlpacaBroker.__new__(AlpacaBroker)
    br._t = fake
    assert br.portfolio_history("1M", "1D") == [
        ("2026-10-09", 1_000_000.0),
        ("2026-10-09", 1_001_000.0),
    ]
    assert br.portfolio_history("1W", "15Min") == [
        ("2026-10-09T00:00", 1_000_000.0),
        ("2026-10-09T14:45", 1_001_000.0),
    ]


def test_cli_from_option_shortens_the_window(tmp_path, capsys):
    out = tmp_path / "o.xlsx"
    assert cp.main(["--demo", "--out", str(out)]) == 0
    full = capsys.readouterr().out
    day = next(ln for ln in full.splitlines() if ln.startswith("Finestra comune")).split()[3]
    import datetime as dt

    later = (dt.date.fromisoformat(day) + dt.timedelta(days=10)).isoformat()
    assert cp.main(["--demo", "--out", str(out), "--from", later]) == 0
    short = capsys.readouterr().out
    assert "Finestra di confronto dal" in short and f"dal {later}" in short.replace(
        "Finestra comune: ", ""
    )
    n_full = int(full.split("rilevazioni")[0].rsplit("(", 1)[1])
    n_short = int(short.split("rilevazioni")[0].rsplit("(", 1)[1])
    assert n_short < n_full
