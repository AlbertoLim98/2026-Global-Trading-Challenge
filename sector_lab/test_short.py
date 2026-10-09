import io
import json
import sys
import threading
import urllib.request
from datetime import datetime
from http.server import ThreadingHTTPServer
from pathlib import Path
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))
import broker as broker_mod
import data
import desk as desk_mod
import journal as journal_mod
import portfolio
import server
import shortterm as st

NY = ZoneInfo("America/New_York")
P = portfolio.Params(mode="short")


def _demo_panel(n=40, seed_syms=None):
    uni = data.load_universe().head(n)
    syms = list(uni.ticker)
    bars = data.demo_bars([*syms, "SPY", *data.SECTORS])
    sector_of = dict(zip(uni.ticker, uni.sector_etf, strict=True))
    sb = {s: bars[s] for s in syms}
    etf = {e: bars[e] for e in data.SECTORS}
    return sb, bars["SPY"], sector_of, etf, uni


def test_panel_has_no_lookahead_truncating_the_future_changes_nothing():
    sb, spy, sector_of, etf, _ = _demo_panel()
    full = st.build_panel(sb, spy, sector_of, etf)
    day = full.close.index[300]
    cut = st.build_panel(
        {k: v.loc[:day] for k, v in sb.items()},
        spy.loc[:day],
        sector_of,
        {k: v.loc[:day] for k, v in etf.items()},
    )
    a, b = full.total.loc[day], cut.total.loc[day]
    assert a.notna().all() and np.allclose(a.to_numpy(), b.to_numpy())
    for p in st.PILLARS:
        assert np.allclose(full.pillars[p].loc[day].to_numpy(), cut.pillars[p].loc[day].to_numpy())


def test_scores_are_percentiles_and_pillars_follow_the_spec():
    sb, spy, sector_of, etf, uni = _demo_panel()
    pn = st.build_panel(sb, spy, sector_of, etf)
    last = pn.close.index[-1]
    assert set(pn.pillars) == set(st.PILLARS) and len(st.PILLARS) == 7
    for p in st.PILLARS:
        x = pn.pillars[p].loc[last].dropna()
        assert x.between(0, 100).all()
    assert pn.total.loc[last].between(0, 100).all()
    # RSI: sotto la media a 50 giorni il pilastro è neutro (50)
    below = pn.close.loc[last] <= pn.sma50.loc[last]
    assert below.any() and (pn.comp_pct["rsi_rev"].loc[last][below] == 50).all()
    rows = st.latest_table(pn, dict(zip(uni.ticker, uni.name, strict=True)), sector_of)
    assert [r["rank"] for r in rows] == list(range(1, len(rows) + 1))
    assert all("stop_mult" not in r for r in rows)  # nessuna regola d'uscita
    assert rows[0]["total"] >= rows[-1]["total"]


def test_incomplete_bar_is_dropped_and_no_stop_helpers_remain():
    assert not hasattr(st, "stop_multiplier") and not hasattr(st, "STOP_BULL")
    df = pd.DataFrame({"close": [1.0, 2.0]}, index=pd.to_datetime(["2026-10-07", "2026-10-08"]))
    before_close = datetime(2026, 10, 8, 11, 0, tzinfo=NY)
    after_close = datetime(2026, 10, 8, 16, 30, tzinfo=NY)
    next_day = datetime(2026, 10, 9, 9, 0, tzinfo=NY)
    assert len(st.drop_incomplete(df, before_close)) == 1  # la barra di oggi è parziale
    assert len(st.drop_incomplete(df, after_close)) == 2
    assert len(st.drop_incomplete(df, next_day)) == 2


def test_ic_report_detects_a_planted_reversal_and_not_noise():
    rng = np.random.default_rng(3)
    n, k = 420, 120
    idx = pd.bdate_range(end="2026-10-08", periods=n)
    spy_r = rng.normal(0.0003, 0.008, n)
    spy = pd.DataFrame({"close": 100 * np.cumprod(1 + spy_r)}, index=idx)
    bars, sector_of = {}, {}
    for j in range(k):
        r = np.zeros(n)
        eps = rng.normal(0, 0.012, n)
        for t in range(1, n):
            r[t] = 0.0002 - 0.5 * eps[t - 1] + eps[t]  # il giorno dopo il movimento di ieri rientra
        c = 50 * np.cumprod(1 + r)
        hi, lo = (
            c * (1 + np.abs(rng.normal(0, 0.004, n))),
            c * (1 - np.abs(rng.normal(0, 0.004, n))),
        )
        bars[f"S{j}"] = pd.DataFrame(
            {
                "open": c,
                "high": hi,
                "low": lo,
                "close": c,
                "volume": rng.integers(1e5, 1e6, n).astype(float),
            },
            index=idx,
        )
        sector_of[f"S{j}"] = "XLK"
    etf = {"XLK": spy}
    rep = st.ic_report(st.build_panel(bars, spy, sector_of, etf))
    by = {r["indicatore"]: r for r in rep["rows"]}
    assert by["componente:move_rev"]["netto_beta"]["ic_t"] > 3  # il segnale piantato viene trovato
    assert by["pilastro:move"]["verdetto"] == "funziona"
    assert abs(by["componente:ctx_beta"]["netto_beta"]["ic_t"]) < 3.5
    assert set(rep["strategy"]) >= {"excess_bps_day", "net_bps_day", "turnover"}


def test_short_targets_top_n_cap_97_percent_without_stops():
    sb, spy, sector_of, etf, uni = _demo_panel(60)
    pn = st.build_panel(sb, spy, sector_of, etf)
    rows = st.latest_table(pn, dict(zip(uni.ticker, uni.name, strict=True)), sector_of)
    tg = portfolio.build_targets_short(rows, P)
    t = tg["targets"]
    assert len(t) == P.short_n == 20 and {r["symbol"] for r in rows[:20]} == set(t)
    vals = [x["value"] for x in t.values()]
    assert sum(vals) == pytest.approx(P.capital * 0.97) and max(vals) <= 0.10 * P.capital + 1e-6
    assert all("stop_mult" not in x for x in t.values())


@pytest.fixture
def short_desk(tmp_path):
    state = server.State(True, "iex")
    syms = [*data.load_universe().ticker, data.STOCK_BENCHMARK, *data.SECTORS, data.BENCHMARK]
    prices = {k: float(v["close"].iloc[-1]) for k, v in data.demo_bars(syms).items()}
    broker = broker_mod.DemoBroker(1_000_000.0, prices)
    j = journal_mod.Journal(tmp_path / "s.db", "Portafoglio 1g", "short")
    d = desk_mod.Desk(
        state, broker, j, settle_seconds=0, portfolio="Portafoglio 1g", strategy="short"
    )
    return d, broker, prices


def test_short_strategy_end_to_end_saves_portfolio_and_strategy_everywhere(short_desk):
    d, broker, _ = short_desk
    res = d.run("quality", 1_000_000)  # la strategia fissa del portafoglio prevale
    assert res["params"]["mode"] == "short" and res["outcome"]["total"] == 20
    assert res["outcome"]["filled"] == 20 and res["n_targets"] == 20
    assert res["portfolio"] == "Portafoglio 1g" and res["strategy"] == "short"
    buys = [t for t in d.trades() if t["kind"] == "BUY"]
    assert len(buys) == 20 and all("stop_mult" not in b for b in buys)
    # su che portafoglio entrano le posizioni e con quale strategia sono calcolate: in ogni operazione
    assert all(b["portfolio"] == "Portafoglio 1g" and b["strategy"] == "short" for b in buys)
    assert all(
        b["strategy_label"] == "Breve termine (1 gg)" and b["account_number"] == "DEMO"
        for b in buys
    )
    run = d.journal.events(kind="RUN")[0]["payload"]
    assert run["portfolio"] == "Portafoglio 1g" and run["strategy"] == "short"
    assert run["strategy_label"] == "Breve termine (1 gg)" and run["account_number"] == "DEMO"
    assert (
        run["short_context"]["state"] in ("Rialzista", "Ribassista")
        and len(run["short_table"]) == 40
    )
    ev = d.journal.events(10_000)
    assert all(e["strategy"] == "short" and e["portfolio"] == "Portafoglio 1g" for e in ev)
    mov = d.portfolio()["movements"]
    assert len(mov) == 20 and all(
        m["portfolio"] == "Portafoglio 1g" and m["strategy"] == "short" for m in mov
    )
    assert d.portfolio()["label"] == "Portafoglio 1g" and d.portfolio()["strategy"] == "short"
    assert len(broker.positions()) == 20


def test_free_instance_saves_the_strategy_used_at_every_run_and_the_account(short_desk, tmp_path):
    d, broker, _ = short_desk
    free = desk_mod.Desk(
        d.state,
        broker_mod.DemoBroker(1_000_000.0, dict(broker.prices)),
        journal_mod.Journal(tmp_path / "free.db"),
        settle_seconds=0,
    )  # istanza senza nome né strategia fissa
    free.run("beta", 1_000_000)
    first = free.journal.events(10_000)
    assert {(e["portfolio"], e["strategy"]) for e in first} == {("Conto DEMO", "beta")}
    free.run("quality", 1_000_000)
    last_run = free.journal.events(kind="RUN")[0]
    assert (last_run["portfolio"], last_run["strategy"]) == ("Conto DEMO", "quality")
    assert last_run["payload"]["strategy_label"] == "Qualità"
    assert [r["strategy"] for r in free.journal.runs()] == ["beta", "quality"]
    assert free.portfolio()["strategy"] == "quality"


def test_short_endpoints_and_excel(short_desk):
    d, _, _ = short_desk
    srv = ThreadingHTTPServer(("127.0.0.1", 0), server.make_handler(d.state, d, 0))
    port = srv.server_address[1]
    srv.server_close()  # libera la porta e riapri con l'elenco di host ammessi corretto
    srv = ThreadingHTTPServer(("127.0.0.1", port), server.make_handler(d.state, d, port))
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:

        def get(path):
            with urllib.request.urlopen(f"http://127.0.0.1:{port}{path}") as r:
                return r.read()

        v = json.loads(get("/api/short"))
        assert v["context"]["state"] in ("Rialzista", "Ribassista") and len(v["rows"]) == 60
        assert set(v["rows"][0]["pillars"]) == set(st.PILLARS) and "_panel" not in v
        ic = json.loads(get("/api/short/ic"))
        assert len(ic["rows"]) == 9 + 7 + 1 and ic["strategy"]["top_n"] == 20
        wb = load_workbook(io.BytesIO(get("/api/short.xlsx?ic=1")))
        assert wb.sheetnames == ["Candidati 1 giorno", "Contesto e note", "Validazione storica"]
        ws = wb["Candidati 1 giorno"]
        heads = [c.value for c in ws[2]]
        assert "Moltiplicatore ATR" not in heads and {"Trend", "Score totale"} <= set(heads)
        assert ws.max_row > 400
        assert load_workbook(io.BytesIO(get("/api/short.xlsx"))).sheetnames == [
            "Candidati 1 giorno",
            "Contesto e note",
        ]
    finally:
        srv.shutdown()
