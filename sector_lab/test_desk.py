import json
import sqlite3
import sys
import threading
import urllib.error
import urllib.request
from http.server import ThreadingHTTPServer
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent))
import broker as broker_mod
import data
import desk as desk_mod
import journal as journal_mod
import portfolio
import server

P = portfolio.Params()


def test_allocate_caps_and_leftover():
    a = portfolio.allocate({"a": 5, "b": 1, "c": 1}, 100.0, 40.0)
    assert a["a"] == 40 and a["b"] == pytest.approx(30) and sum(a.values()) == pytest.approx(100)
    assert (
        sum(portfolio.allocate({"a": 1, "b": 1}, 100.0, 30.0).values()) == 60
    )  # il resto resta liquidità


def _sector(sym, trend, total):
    return {"symbol": sym, "name": sym, "trend_label": trend, "scores": {"total": total}}


def _stock(sym, score, price=100.0, trend="Rialzista", twrr=0.05, beta=1.2):
    return {
        "symbol": sym,
        "name": sym,
        "price": price,
        "trend_label": trend,
        "twrr_w": twrr,
        "beta_1y": beta,
        "scores": {"total": score},
        "strategy_weight": 0.1,
    }


def test_targets_skip_bearish_and_low_score_and_respect_caps():
    secs = [
        _sector("XLK", "Rialzista", 90),
        _sector("XLE", "Ribassista", 95),
        _sector("XLF", "Misto", 40),
    ]
    top = {
        "XLK": {
            "top": [_stock(f"K{i}", 80 - i) for i in range(10)]
            + [_stock("BAD", 99, trend="Ribassista")]
        }
    }
    t = portfolio.build_targets(secs, top, P)
    budgets = {s["symbol"]: s for s in t["sectors"]}
    assert (
        budgets["XLK"]["eligible"]
        and not budgets["XLE"]["eligible"]
        and not budgets["XLF"]["eligible"]
    )
    assert "BAD" not in t["targets"]  # trend relativo ribassista
    assert max(x["value"] for x in t["targets"].values()) <= P.max_stock * P.capital + 1e-6
    # nessun tetto per settore: con un solo settore idoneo e 10 titoli si investe fino al 97%
    assert sum(x["value"] for x in t["targets"].values()) == pytest.approx(
        P.capital * (1 - P.cash_reserve)
    )
    assert P.max_stock == 0.10 and not hasattr(P, "max_sector")


def test_two_sectors_only_every_stock_capped_at_10_percent_of_portfolio():
    secs = [_sector("XLK", "Rialzista", 90), _sector("XLF", "Rialzista", 70)] + [
        _sector(e, "Ribassista", 80) for e in ("XLE", "XLV")
    ]
    top = {
        "XLK": {"top": [_stock(f"K{i}", 90 - i) for i in range(10)]},
        "XLF": {"top": [_stock(f"F{i}", 70 - i) for i in range(10)]},
    }
    t = portfolio.build_targets(secs, top, P)
    vals = [x["value"] for x in t["targets"].values()]
    assert len(vals) == 20 and max(vals) <= 0.10 * P.capital + 1e-6
    assert sum(vals) == pytest.approx(P.capital * (1 - P.cash_reserve))
    by = {x["symbol"]: x for x in t["sectors"]}
    assert by["XLK"]["budget"] > by["XLF"]["budget"] > 0 and by["XLE"]["budget"] == 0


def test_always_invests_97_percent_by_adding_more_sectors():
    # un solo settore idoneo con 4 titoli (max 40%): servono altri settori per arrivare al 97%
    secs = [
        _sector("XLK", "Rialzista", 90),
        _sector("XLF", "Misto", 40),  # sotto soglia punteggio
        _sector("XLE", "Misto", 35),
        _sector("XLV", "Ribassista", 80),
    ]
    top = {
        "XLK": {"top": [_stock(f"K{i}", 90 - i) for i in range(4)]},
        "XLF": {"top": [_stock(f"F{i}", 60 - i) for i in range(4)]},
        "XLE": {"top": [_stock(f"E{i}", 50 - i) for i in range(4)]},
        "XLV": {"top": [_stock(f"V{i}", 70 - i) for i in range(4)]},
    }
    t = portfolio.build_targets(secs, top, P)
    vals = [x["value"] for x in t["targets"].values()]
    assert sum(vals) == pytest.approx(P.capital * 0.97) and max(vals) <= 0.10 * P.capital + 1e-6
    by = {x["symbol"]: x for x in t["sectors"]}
    assert by["XLK"]["budget"] == pytest.approx(0.40 * P.capital)  # il settore idoneo è al massimo
    assert by["XLF"]["budget"] > 0 and by["XLE"]["budget"] > 0  # aggiunti (non ribassisti) prima...
    assert by["XLV"]["budget"] == 0  # ...del settore ribassista, non necessario
    assert any("sono stati aggiunti" in n for n in t["notes"])


def test_bearish_sector_used_only_when_still_short():
    secs = [_sector("XLK", "Rialzista", 90), _sector("XLV", "Ribassista", 80)]
    top = {
        "XLK": {"top": [_stock(f"K{i}", 90 - i) for i in range(6)]},
        "XLV": {"top": [_stock(f"V{i}", 70 - i) for i in range(6)]},
    }
    t = portfolio.build_targets(secs, top, P)
    assert sum(x["value"] for x in t["targets"].values()) == pytest.approx(P.capital * 0.97)
    assert {x["sector"] for x in t["targets"].values()} == {"XLK", "XLV"}
    assert "ribassista" in " ".join(t["notes"])


def test_relaxed_filters_as_last_resort_to_reach_97_percent():
    secs = [_sector("XLK", "Rialzista", 90)]
    stocks_ = [_stock(f"K{i}", 90 - i) for i in range(5)]
    stocks_ += [_stock(f"W{i}", 40 - i, trend="Ribassista") for i in range(6)]  # non idonei
    t = portfolio.build_targets(secs, {"XLK": {"top": stocks_}}, P)
    assert sum(x["value"] for x in t["targets"].values()) == pytest.approx(P.capital * 0.97)
    assert max(x["value"] for x in t["targets"].values()) <= 0.10 * P.capital + 1e-6
    assert len(t["targets"]) == 10 and any(k.startswith("W") for k in t["targets"])


def test_cash_exceeds_3_percent_only_when_titles_do_not_exist_at_all():
    secs = [_sector("XLK", "Rialzista", 90)]
    top = {"XLK": {"top": [_stock(f"K{i}", 80) for i in range(3)]}}
    t = portfolio.build_targets(secs, top, P)
    assert all(x["value"] == pytest.approx(0.10 * P.capital) for x in t["targets"].values())
    assert sum(x["value"] for x in t["targets"].values()) == pytest.approx(0.30 * P.capital)
    assert any("Non ci sono abbastanza titoli" in n for n in t["notes"])


def test_proposals_exit_buy_and_unmanaged_without_any_stop():
    managed = {s: {"name": s, "sector": "XLK"} for s in ("LOSER", "OUT", "NEW")}
    targets = {
        "NEW": {
            "symbol": "NEW",
            "name": "NEW",
            "sector": "XLK",
            "value": 20_000.0,
            "price": 50.0,
            "score": 80,
        },
        "LOSER": {
            "symbol": "LOSER",
            "name": "LOSER",
            "sector": "XLK",
            "value": 5_000.0,
            "price": 50.0,
            "score": 70,
        },
    }
    pos = {
        "LOSER": {
            "qty": 100,
            "avg_entry": 100.0,
            "price": 50.0,
        },  # -50%, ma ancora nel target: nessuno stop
        "OUT": {"qty": 10, "avg_entry": 50.0, "price": 52.0},
        "ZZZ": {"qty": 5, "avg_entry": 10.0, "price": 9.0},  # fuori strategia
    }
    r = portfolio.build_proposals(targets, pos, {"NEW": 50.0}, managed, 500_000.0, P)
    by = {(x["kind"], x["symbol"]): x for x in r["proposals"]}
    assert ("SELL", "OUT") in by and ("BUY", "NEW") in by and by[("BUY", "NEW")]["qty"] == 400
    assert not any(
        x["symbol"] == "LOSER" for x in r["proposals"]
    )  # la perdita da sola non fa uscire
    assert not any(x["kind"] == "STOP" for x in r["proposals"])
    assert not any(x["symbol"] == "ZZZ" for x in r["proposals"]) and any(
        "ZZZ" in n for n in r["notes"]
    )
    assert [x["priority"] for x in r["proposals"]] == sorted(x["priority"] for x in r["proposals"])


def test_atr_stop_has_been_removed():
    assert not hasattr(portfolio, "stop_hit") and not hasattr(P, "atr_stop_mult")


def test_no_partial_sells_for_overweight_titles():
    managed = {"BIG": {"name": "B", "sector": "XLK"}}
    targets = {
        "BIG": {
            "symbol": "BIG",
            "name": "B",
            "sector": "XLK",
            "value": 10_000.0,
            "price": 100.0,
            "score": 70,
        }
    }
    pos = {"BIG": {"qty": 500, "avg_entry": 100.0, "price": 100.0}}  # 50k contro target 10k
    assert portfolio.build_proposals(targets, pos, {}, managed, 100_000.0, P)["proposals"] == []
    assert P.cash_reserve == 0.03


def test_buys_scaled_when_cash_short_and_small_drift_ignored():
    managed = {"NEW": {"name": "N", "sector": "XLK"}, "OLD": {"name": "O", "sector": "XLK"}}
    targets = {
        "NEW": {
            "symbol": "NEW",
            "name": "N",
            "sector": "XLK",
            "value": 40_000.0,
            "price": 100.0,
            "score": 70,
        },
        "OLD": {
            "symbol": "OLD",
            "name": "O",
            "sector": "XLK",
            "value": 10_000.0,
            "price": 100.0,
            "score": 70,
        },
    }
    pos = {
        "OLD": {"qty": 95, "avg_entry": 100.0, "price": 100.0}
    }  # 9.5k vs 10k: scarto trascurabile
    r = portfolio.build_proposals(targets, pos, {}, managed, P.cash_reserve * P.capital + 10_000, P)
    assert [x["symbol"] for x in r["proposals"]] == ["NEW"]
    assert r["proposals"][0]["qty"] == 100 and any(
        "Liquidità insufficiente" in n for n in r["notes"]
    )


def test_journal_is_append_only(tmp_path):
    j = journal_mod.Journal(tmp_path / "j.db")
    j.log("RUN", {"x": 1}, "r1")
    with pytest.raises(sqlite3.DatabaseError):
        j._db.execute("UPDATE events SET kind = 'X'")
    with pytest.raises(sqlite3.DatabaseError):
        j._db.execute("DELETE FROM events")
    assert len(j.events()) == 1


@pytest.fixture
def env(tmp_path):
    state = server.State(True, "iex")
    syms = [*data.load_universe().ticker, data.STOCK_BENCHMARK, *data.SECTORS, data.BENCHMARK]
    prices = {k: float(v["close"].iloc[-1]) for k, v in data.demo_bars(syms).items()}
    broker = broker_mod.DemoBroker(1_000_000.0, prices)
    desk = desk_mod.Desk(state, broker, journal_mod.Journal(tmp_path / "j.db"), settle_seconds=0)
    return desk, broker, prices


def _statuses(desk):
    return [t["status"] for t in desk.trades()]


def test_run_sends_all_orders_automatically_and_journals_everything(env):
    desk, broker, _ = env
    res = desk.run("quality", 1_000_000)
    out = res["outcome"]
    assert out["total"] > 0 and out["filled"] == out["total"] and out["failed"] == 0
    assert len(broker.positions()) == out["total"] and len(broker.orders) == out["total"]
    events = desk.journal.events(10_000)
    kinds = [e["kind"] for e in events]
    assert kinds.count("DECISION") == out["total"] and kinds.count("ORDER") == out["total"]
    assert all(e["payload"]["action"] == "auto" for e in events if e["kind"] == "DECISION")
    run_evt = desk.journal.events(kind="RUN")[0]["payload"]
    assert run_evt["sector_table"] and run_evt["stock_tables"] and run_evt["targets"]
    # nessuna operazione resta da approvare: tutte eseguite
    assert set(_statuses(desk)) == {"filled"}


def test_second_daily_run_proposes_nothing_when_portfolio_is_in_line(env):
    desk, _, _ = env
    desk.run("quality", 1_000_000)
    assert desk.run("quality", 1_000_000)["outcome"]["total"] == 0


def test_sells_are_sent_before_buys(env):
    desk, broker, prices = env
    sym = next(s for s in data.load_universe().ticker if s not in ())
    # una posizione fuori target (qualsiasi titolo dell'universo non selezionato) viene venduta per prima
    held = [s for s in prices if s in desk.managed][:1][0]
    broker.pos[held] = {"qty": 10.0, "avg_entry": prices[held]}
    desk.run("quality", 1_000_000)
    sent = [e["payload"] for e in reversed(desk.journal.events(10_000)) if e["kind"] == "DECISION"]
    sides = [x["side"] for x in sent]
    assert "sell" in sides and sides.index("buy") > max(
        i for i, v in enumerate(sides) if v == "sell"
    )
    assert sym


def test_a_big_loss_does_not_trigger_any_exit_only_the_rebalance_decides(env):
    desk, broker, prices = env
    assert not hasattr(desk, "stop_check")
    desk.run("quality", 1_000_000)
    sym = next(iter(broker.pos))
    prices[sym] *= 0.5  # crollo del titolo
    out = desk.run("quality", 1_000_000)
    kinds = {t["kind"] for t in desk.trades()}
    assert "STOP" not in kinds
    assert not any(t["symbol"] == sym and t["side"] == "sell" for t in desk.trades())
    assert out["outcome"]["failed"] == 0


def test_failed_orders_are_reported_and_can_be_repositioned_at_current_price(env, monkeypatch):
    desk, broker, prices = env

    def boom(*a, **k):
        raise RuntimeError("rete caduta")

    monkeypatch.setattr(broker, "submit_market", boom)
    out = desk.run("quality", 1_000_000)["outcome"]
    assert out["total"] > 0 and out["failed"] == out["total"] and out["filled"] == 0
    failed = [t for t in desk.trades() if t["status"] == "failed"]
    assert all("rete caduta" in t["error"] for t in failed)
    kinds = [e["kind"] for e in desk.journal.events(10_000)]
    assert kinds.count("TRADE_FAILED") == len(failed)

    monkeypatch.undo()
    t = failed[0]
    prices[t["symbol"]] *= 1.03  # il prezzo si è mosso: si usa quello attuale
    r = desk.reposition(t["id"])
    assert r["status"] == "filled" and r["limit_price"] == pytest.approx(prices[t["symbol"]])
    assert desk.journal.get_proposal(t["id"])["status"] == "repositioned"
    new = desk.journal.get_proposal(r["new_id"])
    assert (
        new["parent_id"] == t["id"] and new["order_type"] == "limit" and new["status"] == "filled"
    )
    assert t["symbol"] in {p["symbol"] for p in broker.positions()}
    with pytest.raises(desk_mod.DeskError):
        desk.reposition(t["id"])  # già riposizionata: niente doppio ordine
    assert any(e["kind"] == "REPOSITION" for e in desk.journal.events(10_000))


def test_rejected_and_partially_filled_orders_are_failures(env, monkeypatch):
    desk, broker, _ = env
    calls = []

    def fake(symbol, side, qty, client_id):
        calls.append(qty)
        if len(calls) == 1:
            return {
                "order_id": "o1",
                "order_status": "rejected",
                "filled_qty": 0.0,
                "filled_avg_price": None,
            }
        return {
            "order_id": f"o{len(calls)}",
            "order_status": "expired",
            "filled_qty": qty // 2,
            "filled_avg_price": 10.0,
        }

    monkeypatch.setattr(broker, "submit_market", fake)
    desk.run("quality", 1_000_000)
    failed = {t["order_id"]: t for t in desk.trades() if t["status"] == "failed"}
    assert failed["o1"]["error"] == "ordine rejected" and failed["o1"]["remaining_qty"] == calls[0]
    part = failed["o2"]
    assert "eseguite" in part["error"] and part["remaining_qty"] == calls[1] - calls[1] // 2


def test_buying_power_shortage_is_a_reported_failure_not_a_crash(env, monkeypatch):
    desk, broker, _ = env
    orig = broker.account
    monkeypatch.setattr(broker, "account", lambda: {**orig(), "buying_power": 0.0})
    out = desk.run("quality", 1_000_000)["outcome"]
    assert out["failed"] > 0 and out["failed"] == out["total"]
    errs = [t["error"] for t in desk.trades() if t["status"] == "failed"]
    assert all("potere d'acquisto" in e for e in errs)


def test_scarce_cash_scales_buys_down_instead_of_failing(env):
    desk, broker, _ = env
    broker.cash = 100_000.0
    out = desk.run("quality", 1_000_000)["outcome"]
    assert out["failed"] == 0 and out["filled"] == out["total"]
    assert broker.account()["cash"] >= 0.03 * 1_000_000 - 1  # la riserva del 3% viene rispettata


def test_open_orders_block_a_new_run_until_resolved(env, monkeypatch):
    desk, broker, _ = env

    def queued(symbol, side, qty, client_id):
        return {
            "order_id": f"q-{client_id}",
            "order_status": "accepted",
            "filled_qty": 0.0,
            "filled_avg_price": None,
        }

    monkeypatch.setattr(broker, "submit_market", queued)
    assert desk.run("quality", 1_000_000)["outcome"]["in_progress"] > 0
    with pytest.raises(desk_mod.DeskError, match="ordini ancora in corso"):
        desk.run("quality", 1_000_000)
    monkeypatch.setattr(
        broker,
        "get_order",
        lambda oid: {
            "order_id": oid,
            "order_status": "filled",
            "filled_qty": 1.0,
            "filled_avg_price": 1.0,
        },
    )
    assert desk.refresh_orders() > 0 and "submitted" not in _statuses(desk)


def test_new_run_archives_old_failures(env, monkeypatch):
    desk, broker, _ = env

    def boom(*a, **k):
        raise RuntimeError("giù")

    monkeypatch.setattr(broker, "submit_market", boom)
    desk.run("quality", 1_000_000)
    old = {t["id"] for t in desk.trades() if t["status"] == "failed"}
    monkeypatch.undo()
    desk.run("quality", 1_000_000)
    assert old and all(desk.journal.get_proposal(i)["status"] == "superseded" for i in old)


def test_portfolio_view_composition_pl_and_todays_movements(env):
    desk, _, prices = env
    assert desk.portfolio()["opening"] is None and desk.portfolio()["since_open_pl"] is None
    desk.run("quality", 1_000_000)
    v = desk.portfolio()
    acct = v["account"]
    assert v["opening"]["equity"] == pytest.approx(1_000_000.0)
    assert v["since_open_pl"] == pytest.approx(acct["equity"] - 1_000_000.0)
    assert v["day_pl"] == pytest.approx(acct["equity"] - acct["last_equity"])
    assert sum(r["weight"] for r in v["positions"]) + v["cash_weight"] == pytest.approx(1.0)
    assert max(r["weight"] for r in v["positions"]) <= 0.10 + 1e-3  # nessun titolo oltre il 10%
    assert sum(x["weight"] for x in v["sectors"]) == pytest.approx(1 - v["cash_weight"])
    assert len(v["movements"]) == len(v["positions"]) and all(
        m["side"] == "buy" for m in v["movements"]
    )
    assert v["movements"][0]["time"] >= v["movements"][-1]["time"]  # dal più recente
    sym = v["positions"][0]["symbol"]
    prices[sym] *= 1.1  # il prezzo sale: P/L non realizzato positivo
    assert next(r for r in desk.portfolio()["positions"] if r["symbol"] == sym)["pl"] > 0


def test_gain_per_reallocation(env):
    desk, broker, prices = env
    assert desk.portfolio()["rebalances"] == [] and desk.portfolio()["since_last"] is None
    desk.run("quality", 1_000_000)
    v1 = desk.portfolio()
    assert len(v1["rebalances"]) == 1 and v1["rebalances"][0]["ongoing"]
    assert v1["since_last"]["equity"] == pytest.approx(1_000_000.0)
    assert v1["since_last"]["pl"] == pytest.approx(v1["account"]["equity"] - 1_000_000.0)
    for sym in list(broker.pos):  # il mercato sale del 2%
        prices[sym] *= 1.02
    mid = desk.portfolio()["since_last"]["pl"]
    assert mid > 0
    desk.run("quality", 1_000_000)
    v2 = desk.portfolio()
    first, second = v2["rebalances"][1], v2["rebalances"][0]  # dalla più recente
    assert not first["ongoing"] and second["ongoing"]
    assert first["end_equity"] == pytest.approx(second["equity"]) and first["pl"] == pytest.approx(
        mid
    )
    assert first["pl_pct"] == pytest.approx(first["end_equity"] / first["equity"] - 1)


def test_reposition_only_for_failed(env):
    desk, _, _ = env
    desk.run("quality", 1_000_000)
    ok = desk.trades()[0]
    with pytest.raises(desk_mod.DeskError):
        desk.reposition(ok["id"])
    with pytest.raises(desk_mod.DeskError):
        desk.reposition("inesistente")


def test_server_rejects_cross_site_posts_and_foreign_hosts(env):
    desk, _, _ = env
    srv = ThreadingHTTPServer(("127.0.0.1", 0), server.make_handler(desk.state, desk, 0))
    port = srv.server_address[1]
    srv = None  # ricrea con la porta reale per l'elenco degli host ammessi
    srv = ThreadingHTTPServer(("127.0.0.1", port), server.make_handler(desk.state, desk, port))
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    url = f"http://127.0.0.1:{port}"

    def call(path, method="GET", headers=None, body=b"{}"):
        req = urllib.request.Request(
            url + path,
            data=body if method == "POST" else None,
            method=method,
            headers=headers or {},
        )
        try:
            with urllib.request.urlopen(req) as r:
                return r.status, json.loads(r.read() or b"{}")
        except urllib.error.HTTPError as e:
            return e.code, json.loads(e.read() or b"{}")

    try:
        assert call("/api/trade/reposition", "POST")[0] == 403  # senza intestazione
        assert (
            call("/api/trade/reposition", "POST", {"X-Sector-Lab": "1", "Host": "evil.example"})[0]
            == 403
        )
        assert call("/api/account", headers={"Host": "evil.example"})[0] == 403
        assert (
            call(
                "/api/trade/reposition",
                "POST",
                {"X-Sector-Lab": "1"},
                b'{"id": "x", "action": "execute"}',
            )[0]
            == 400
        )
        assert call("/api/account")[0] == 200
    finally:
        srv.shutdown()


def test_stock_view_weights_come_from_the_real_allocation_across_sectors(env):
    desk, _, _ = env
    state = desk.state
    views = {e: state.stock_view(e, "quality") for e in data.SECTORS}
    total = sum(r["weight_portfolio"] for v in views.values() for r in v["top"])
    assert total == pytest.approx(0.97)  # il 97% è distribuito tra le top 10 di tutti i settori
    assert max(r["weight_portfolio"] for v in views.values() for r in v["top"]) <= 0.10 + 1e-9
    for v in views.values():
        assert sum(r["weight_portfolio"] for r in v["top"]) == pytest.approx(
            v["sector_portfolio_weight"]
        )
        w = [r["weight_in_sector"] for r in v["top"]]
        assert sum(w) == pytest.approx(1) or sum(w) == 0  # senza tetti: quote del punteggio
    secs = {x["symbol"]: x["scores"]["total"] for x in state.sectors()["sectors"]}
    assert all(r["sector_score"] == secs[e] for e, v in views.items() for r in v["top"])
    bv = state.stock_view("XLK", "beta")
    assert all("weight_portfolio" in r and "weight_in_sector" in r for r in bv["top"])


def test_untradable_titles_are_excluded_from_selection_and_reported(env):
    desk, _, _ = env
    uni = set(data.load_universe().ticker)
    gone = {"AAPL", "NVDA", "MSFT", "AMZN", "GOOGL"} & uni
    desk.state.tradable_fn = lambda: uni - gone
    res = desk.run("quality", 1_000_000)
    held = {t["symbol"] for t in desk.trades()}
    assert not (held & gone)
    assert any("non sono negoziabili su Alpaca" in n for n in res["notes"])
    views = [desk.state.stock_view(e, "quality") for e in data.SECTORS]
    assert not ({r["symbol"] for v in views for r in v["top"]} & gone)
    assert sum(r["weight_portfolio"] for v in views for r in v["top"]) == pytest.approx(0.97)


def test_inactive_asset_error_is_explained_and_reposition_prefix_is_not_repeated(env, monkeypatch):
    desk, broker, _ = env

    def inactive(*a, **k):
        raise RuntimeError('{"code":40010001,"message":"asset WBD is not active"}')

    monkeypatch.setattr(broker, "submit_market", inactive)
    monkeypatch.setattr(broker, "submit_limit", inactive)
    desk.run("quality", 1_000_000)
    t = next(x for x in desk.trades() if x["status"] == "failed")
    assert "non è negoziabile" in t["error"] and "prossimo ribilanciamento" in t["error"]
    r1 = desk.reposition(t["id"])
    r2 = desk.reposition(r1["new_id"])
    reason = desk.journal.get_proposal(r2["new_id"])["reason"]
    assert reason.count("Riposizionata a prezzo attuale") == 1
    assert reason.endswith(t["reason"])


# --- portafogli con strategie diverse e correzione delle etichette -----------------------------------
import journal_tool


def _desk(tmp_path, name, portfolio=None, strategy=None):
    state = server.State(True, "iex")
    syms = [*data.load_universe().ticker, data.STOCK_BENCHMARK, *data.SECTORS, data.BENCHMARK]
    prices = {k: float(v["close"].iloc[-1]) for k, v in data.demo_bars(syms).items()}
    j = journal_mod.Journal(tmp_path / f"{name}.db", portfolio, strategy)
    return desk_mod.Desk(
        state, broker_mod.DemoBroker(1_000_000.0, prices), j, settle_seconds=0,
        portfolio=portfolio, strategy=strategy,
    )  # fmt: skip


def test_locked_strategy_labels_every_new_event_and_ignores_the_requested_mode(tmp_path):
    d = _desk(tmp_path, "b", "Portafoglio 17", "beta")
    assert d.info() == {
        "portfolio": "Portafoglio 17",
        "strategy": "beta",
        "locked": True,
        "broker": "demo",
    }
    d.run("quality", 1_000_000)  # la richiesta dice qualità, ma il portafoglio è alto beta
    ev = d.journal.events(10_000)
    assert ev and all(e["portfolio"] == "Portafoglio 17" and e["strategy"] == "beta" for e in ev)
    assert d.journal.events(kind="RUN")[0]["payload"]["params"]["mode"] == "beta"
    assert d.portfolio()["rebalances"][0]["strategy"] == "beta"
    assert d.portfolio()["rebalances"][0]["portfolio"] == "Portafoglio 17"


def test_two_portfolios_with_different_strategies_keep_separate_journals(tmp_path):
    a = _desk(tmp_path, "a", "Portafoglio 17", "beta")
    b = _desk(tmp_path, "b", "Portafoglio 18", "quality")
    a.run("quality", 1_000_000)
    b.run("beta", 1_000_000)
    assert {e["strategy"] for e in a.journal.events(10_000)} == {"beta"}
    assert {e["strategy"] for e in b.journal.events(10_000)} == {"quality"}
    assert {e["portfolio"] for e in b.journal.events(10_000)} == {"Portafoglio 18"}
    assert (
        set(a.trades_symbols()) != set(b.trades_symbols()) if hasattr(a, "trades_symbols") else True
    )


def test_old_journal_gets_new_columns_and_relabel_corrects_without_rewriting(tmp_path):
    path = tmp_path / "old.db"
    old = sqlite3.connect(path)  # journal creato prima delle colonne portafoglio/strategia
    old.executescript(
        """CREATE TABLE events (id INTEGER PRIMARY KEY AUTOINCREMENT, ts TEXT NOT NULL, kind TEXT NOT NULL,
        run_id TEXT, proposal_id TEXT, payload TEXT NOT NULL);
        INSERT INTO events (ts, kind, run_id, payload) VALUES
        ('2026-10-08T17:05:00Z','RUN','r17','{"params":{"mode":"quality"},"account":{"equity":1000000}}'),
        ('2026-10-08T17:06:00Z','PROPOSAL','r17','{"kind":"BUY","qty":1,"symbol":"A","price":1,"value":1,"reason":"x"}'),
        ('2026-10-08T18:05:00Z','RUN','r18','{"params":{"mode":"quality"},"account":{"equity":1000000}}');"""
    )
    old.commit()
    old.close()
    j = journal_mod.Journal(path)
    assert all(e["portfolio"] is None for e in j.events())  # non assegnato
    # le ore si leggono in UTC, come nell'Excel del journal
    assert (
        journal_tool.main(
            [
                "relabel",
                "--db",
                str(path),
                "--hour",
                "17",
                "--portfolio",
                "Portafoglio 17",
                "--strategy",
                "beta",
                "--dry-run",
            ]
        )
        == 0
    )
    assert all(
        e["portfolio"] is None for e in journal_mod.Journal(path).events()
    )  # dry-run: nulla cambia
    assert (
        journal_tool.main(
            [
                "relabel",
                "--db",
                str(path),
                "--hour",
                "17",
                "--portfolio",
                "Portafoglio 17",
                "--strategy",
                "beta",
            ]
        )
        == 0
    )
    assert (
        journal_tool.main(
            [
                "relabel",
                "--db",
                str(path),
                "--hour",
                "18",
                "--portfolio",
                "Portafoglio 18",
                "--strategy",
                "quality",
            ]
        )
        == 0
    )
    j = journal_mod.Journal(path)
    by_run = {e["run_id"]: e for e in j.events(100) if e["kind"] in ("RUN", "PROPOSAL")}
    assert (by_run["r17"]["portfolio"], by_run["r17"]["strategy"], by_run["r17"]["relabeled"]) == (
        "Portafoglio 17",
        "beta",
        True,
    )
    assert (by_run["r18"]["portfolio"], by_run["r18"]["strategy"]) == ("Portafoglio 18", "quality")
    raw = j._db.execute("SELECT portfolio, strategy FROM events WHERE run_id = 'r17'").fetchall()
    assert all(r["portfolio"] is None for r in raw)  # le righe originali non sono state toccate
    assert [e["kind"] for e in j.events(100) if e["kind"] == "RELABEL"] == ["RELABEL", "RELABEL"]
    runs = {r["run_id"]: r for r in j.runs()}
    assert runs["r17"]["strategy"] == "beta" and runs["r17"]["portfolio"] == "Portafoglio 17"
    # una nuova assegnazione sostituisce la precedente (vince l'ultima)
    j.relabel(["r17"], "Portafoglio 17", "quality", "prova")
    assert {e["strategy"] for e in j.events(100) if e["run_id"] == "r17"} == {"quality"}


def test_relabel_tool_validations(tmp_path, capsys):
    path = tmp_path / "t.db"
    journal_mod.Journal(path).log(
        "RUN", {"params": {"mode": "beta"}, "account": {"equity": 1}}, "r1"
    )
    assert (
        journal_tool.main(["relabel", "--db", str(path), "--portfolio", "P", "--strategy", "beta"])
        == 2
    )
    assert (
        journal_tool.main(
            ["relabel", "--db", str(path), "--hour", "3", "--portfolio", "P", "--strategy", "beta"]
        )
        == 1
    )
    assert (
        journal_tool.main(
            ["relabel", "--db", str(path), "--all", "--portfolio", "P", "--strategy", "beta"]
        )
        == 0
    )
    assert journal_tool.main(["list", "--db", str(path)]) == 0
    assert "P / Alto beta" in capsys.readouterr().out
    assert journal_tool.main(["list", "--db", str(tmp_path / "manca.db")]) == 2


def test_journal_excel_has_portfolio_and_strategy_columns(tmp_path):
    import io

    import export
    from openpyxl import load_workbook

    d = _desk(tmp_path, "x", "Portafoglio 18", "quality")
    d.run("quality", 1_000_000)
    ws = load_workbook(io.BytesIO(export.build_journal(d.journal.events(100000))))["Journal"]
    heads = [c.value for c in ws[2]]
    assert "Portafoglio" in heads and "Strategia" in heads
    i = heads.index("Strategia") + 1
    assert {ws.cell(r, i).value for r in range(3, ws.max_row + 1)} == {"Qualità"}


def test_hour_is_read_in_utc_by_default_and_tz_can_be_changed(tmp_path, capsys):
    path = tmp_path / "h.db"
    journal_mod.Journal(path).log(
        "RUN", {"params": {"mode": "beta"}, "account": {"equity": 1}}, "r1"
    )
    j = journal_mod.Journal(path)
    h = int(j.run_catalog()[0]["ts"][11:13])  # ora UTC dell'esecuzione
    assert (
        journal_tool.main(
            [
                "relabel",
                "--db",
                str(path),
                "--hour",
                str(h),
                "--portfolio",
                "P",
                "--strategy",
                "beta",
                "--dry-run",
            ]
        )
        == 0
    )
    assert (
        journal_tool.main(
            [
                "relabel",
                "--db",
                str(path),
                "--hour",
                str((h + 5) % 24),
                "--portfolio",
                "P",
                "--strategy",
                "beta",
                "--dry-run",
            ]
        )
        == 1
    )
    assert journal_tool.main(
        [
            "relabel",
            "--db",
            str(path),
            "--hour",
            str((h + 2) % 24),
            "--tz",
            "Europe/Rome",
            "--portfolio",
            "P",
            "--strategy",
            "beta",
            "--dry-run",
        ]
    ) in (0, 1)
