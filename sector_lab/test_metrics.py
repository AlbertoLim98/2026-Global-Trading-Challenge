import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).parent))
import data
import metrics


def frame(close, volume=None):
    idx = pd.bdate_range(end="2026-10-06", periods=len(close))
    close = np.asarray(close, float)
    vol = np.full(len(close), 1e6) if volume is None else np.asarray(volume, float)
    return pd.DataFrame(
        {"open": close, "high": close * 1.01, "low": close * 0.99, "close": close, "volume": vol},
        index=idx,
    )


def test_uptrend_label_and_returns():
    df = frame(100 * 1.001 ** np.arange(300))
    m = metrics.compute(df)
    assert m["trend_label"] == "Rialzista"
    assert m["ret_3m"] == pytest.approx(1.001**63 - 1)
    assert m["drawdown_52w"] == pytest.approx(0)
    assert m["rsi14"] == 100


def test_downtrend_label():
    assert metrics.compute(frame(100 * 0.999 ** np.arange(300)))["trend_label"] == "Ribassista"


def test_short_history_gives_none_not_crash():
    m = metrics.compute(frame(100 + np.arange(60.0)))
    assert m["px_vs_sma200"] is None and m["ret_12m"] is None and m["trend_label"] == "n/d"


def test_volume_ratio_and_updown():
    vol = np.r_[np.full(280, 1e6), np.full(20, 2e6)]
    close = np.r_[np.full(280, 100.0), 100 + np.arange(1.0, 21)]
    m = metrics.compute(frame(close, vol))
    assert m["vol_ratio_20_90"] == pytest.approx(2e6 / ((70 * 1e6 + 20 * 2e6) / 90))
    assert m["updown_volume"] is None  # nessun giorno in calo -> indefinito


def test_relative_strength_vs_benchmark():
    s = frame(100 * 1.002 ** np.arange(300))
    b = frame(100 * 1.001 ** np.arange(300))["close"]
    m = metrics.compute(s, b)
    assert m["rs_3m"] == pytest.approx((1.002**63 - 1) - (1.001**63 - 1))


def test_scores_rank_and_low_vol_is_better():
    rows = {}
    for sym, drift, sig in [("A", 0.002, 0.005), ("B", 0.0, 0.01), ("C", -0.001, 0.02)]:
        rng = np.random.default_rng(1)
        rows[sym] = metrics.compute(frame(100 * np.exp(np.cumsum(rng.normal(drift, sig, 300)))))
    sc = metrics.score(rows)
    assert sc["A"]["momentum"] > sc["B"]["momentum"] > sc["C"]["momentum"]
    assert sc["A"]["volatility"] > sc["C"]["volatility"]
    assert all(0 <= v <= 100 for s in sc.values() for v in s.values() if v is not None)


def test_history_has_no_nan():
    h = metrics.history(frame(100 + np.arange(500.0)))
    assert len(h["dates"]) == 252 and None not in h["close"] and None not in h["sma200"]


def test_demo_bars_cover_all_sectors_and_json_safe():
    import json

    bars = data.demo_bars([*data.SECTORS, data.BENCHMARK])
    rows = {s: metrics.compute(bars[s], bars["SPY"]["close"]) for s in data.SECTORS}
    json.dumps({"r": rows, "s": metrics.score(rows)}, allow_nan=False)


def test_xlsx_export_roundtrip():
    import io

    import export
    from openpyxl import load_workbook

    bars = data.demo_bars([*data.SECTORS, data.BENCHMARK])
    rows = {s: metrics.compute(bars[s], bars["SPY"]["close"]) for s in data.SECTORS}
    sc = metrics.score(rows)
    payload = {
        "demo": True,
        "feed": "demo",
        "updated": "2026-10-08 10:00:00",
        "sectors": [
            {"symbol": s, "name": data.SECTORS[s], **rows[s], "scores": sc[s]} for s in rows
        ],
    }
    ws = load_workbook(io.BytesIO(export.build(payload)))["Settori"]
    assert ws.max_row == 2 + len(data.SECTORS)
    totals = [ws.cell(r, ws.max_column).value for r in range(3, ws.max_row + 1)]
    assert totals == sorted(totals, reverse=True)


def _stock_universe(n=12):
    import pandas as pd

    bars = data.demo_bars([f"T{i}" for i in range(n)] + ["ACWI"])
    uni = pd.DataFrame({"ticker": [f"T{i}" for i in range(n)], "name": "x", "weight_pct": 0.1})
    return uni, bars


def test_twrr_is_ratio_of_linked_returns():
    import stocks

    s = frame(100 * 1.002 ** np.arange(300))
    b = frame(100 * 1.001 ** np.arange(300))
    m = stocks.compute_relative(s, b)
    assert m["twrr_3m"] == pytest.approx(1.002**63 / 1.001**63 - 1)
    exp = sum(w * (1.002**n / 1.001**n - 1) for n, w in stocks.TWRR_WINDOWS.items()) / 10
    assert m["twrr_w"] == pytest.approx(exp)
    assert m["trend_label"] == "Rialzista"


def test_short_history_excluded_and_top10_sorted():
    import stocks

    uni, bars = _stock_universe(14)
    bars["T0"] = bars["T0"].tail(100)  # IPO recente
    r = stocks.rank_sector(uni, bars, bars["ACWI"])
    assert len(r["top"]) == 10 and r["n_skipped"] == 1
    assert "T0" not in [t["symbol"] for t in r["top"]]
    tot = [t["scores"]["total"] for t in r["top"]]
    assert tot == sorted(tot, reverse=True) and [t["rank"] for t in r["top"]] == list(range(1, 11))


def test_stock_export_has_sheet_per_sector():
    import io
    import json

    import export
    import stocks
    from openpyxl import load_workbook

    uni, bars = _stock_universe(12)
    res = stocks.rank_sector(uni, bars, bars["ACWI"])
    json.dumps(res, allow_nan=False)
    meta = {"demo": True, "feed": "demo", "updated": "x"}
    wb = load_workbook(
        io.BytesIO(export.build_stocks({"XLK": res, "XLE": res}, data.SECTORS, meta))
    )
    assert (
        wb.sheetnames[0] == "Tutti"
        and "XLK Tecnologia" in wb.sheetnames
        and "Note" in wb.sheetnames
    )
    assert wb["XLK Tecnologia"].max_row == 12


def test_universe_csv_covers_all_sectors():
    uni = data.load_universe()
    assert set(uni.sector_etf) == set(data.SECTORS) and uni.ticker.is_unique


def test_beta_of_scaled_benchmark_returns():
    import stocks

    rng = np.random.default_rng(7)
    br = rng.normal(0.0003, 0.01, 300)
    b = frame(100 * np.cumprod(1 + br))
    s = frame(100 * np.cumprod(1 + 1.8 * br))
    m = stocks.compute_relative(s, b)
    assert m["beta_1y"] == pytest.approx(1.8, abs=0.02) and m["corr_1y"] == pytest.approx(
        1, abs=1e-6
    )


def test_score_shares_follow_score_with_no_cap():
    import stocks

    rows = [{"scores": {"total": t}} for t in (90.0, 5.0, 5.0)]
    w = stocks.score_shares(rows)
    assert w == pytest.approx([0.9, 0.05, 0.05]) and not hasattr(
        stocks, "MAX_WEIGHT"
    )  # nessun tetto al 25%
    assert stocks.score_shares([{"scores": {"total": None}}] * 4) == pytest.approx([0.25] * 4)
    assert stocks.score_shares([]) == []


def test_beta_mode_filters_and_sorts_by_beta():
    import stocks

    uni, bars = _stock_universe(30)
    rows = stocks.analyze_sector(uni, bars, bars["ACWI"])
    r = stocks.top_beta(rows, len(uni))
    betas = [t["beta_1y"] for t in r["top"]]
    assert betas == sorted(betas, reverse=True) and len(r["top"]) <= 10
    assert all(stocks.is_eligible(t) for t in r["top"])
    assert sum(t["weight_in_sector"] for t in r["top"]) == pytest.approx(1)
    tot = sum(t["scores"]["total"] for t in r["top"])
    assert all(t["weight_in_sector"] == pytest.approx(t["scores"]["total"] / tot) for t in r["top"])
    assert r["n_eligible"] == sum(stocks.is_eligible(x) for x in rows.values())
    assert r["portfolio_beta"] == pytest.approx(
        sum(t["beta_1y"] * t["weight_in_sector"] for t in r["top"])
    )


def test_beta_export_has_weight_column():
    import io

    import export
    import stocks
    from openpyxl import load_workbook

    uni, bars = _stock_universe(30)
    res = stocks.rank_sector(uni, bars, bars["ACWI"], mode="beta")
    meta = {"demo": True, "feed": "demo", "updated": "x"}
    ws = load_workbook(io.BytesIO(export.build_stocks({"XLK": res}, data.SECTORS, meta, "beta")))[
        "XLK Tecnologia"
    ]
    heads = [c.value for c in ws[2]]
    assert {"Peso nel settore", "Peso nel portafoglio", "Beta 1 anno"} <= set(heads)
